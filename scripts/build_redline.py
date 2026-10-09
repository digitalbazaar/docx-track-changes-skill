#!/usr/bin/env python3
# Copyright 2026 Digital Bazaar, Inc.
#
# SPDX-License-Identifier: BSD-3-Clause
"""Diff two Markdown files and emit a .docx with real Word tracked-changes
(w:ins/w:del OOXML), not colored/struck text pretending to be track changes.
Word and Google Docs (on import) both render these as genuine accept/reject
suggestions.

Supports: headings (#..######), '- ' bullet lists (one level of nesting),
pipe tables, **bold**, _italic_/*italic*, ~strike~/~~strike~~, `code`, and
blank-line-separated paragraphs. This is a deliberately narrow Markdown
subset, not a CommonMark engine -- see README.md for exact scope and the
block-vs-word diffing strategy.
"""
import argparse
import itertools
import os
import re
from datetime import datetime, timezone
from difflib import SequenceMatcher

DEFAULT_AUTHOR = os.environ.get("REDLINE_AUTHOR")

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

_rev_id = itertools.count(1)


# ---------- block parsing ----------

def parse_blocks(text):
    lines = text.split("\n")
    blocks = []
    i = 0
    n = len(lines)

    def is_table_line(l):
        return l.strip().startswith("|")

    def is_sep_row(l):
        return re.match(r"^\s*\|[\s:\-\|]+\|\s*$", l) is not None

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped == "":
            i += 1
            continue

        if stripped == "---":
            blocks.append({"type": "hr", "text": "---"})
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            blocks.append({"type": "h", "level": len(m.group(1)), "text": m.group(2).strip()})
            i += 1
            continue

        if is_table_line(line):
            rows = []
            while i < n and is_table_line(lines[i]):
                if not is_sep_row(lines[i]):
                    raw = lines[i].strip()
                    cells = [c.strip() for c in raw.strip("|").split("|")]
                    rows.append(cells)
                i += 1
            blocks.append({"type": "table", "rows": rows,
                            "text": "\n".join("|".join(r) for r in rows)})
            continue

        m = re.match(r"^(\s*)-\s+(.*)$", line)
        if m:
            depth = len(m.group(1)) // 2
            buf = [m.group(2).rstrip()]
            i += 1
            while i < n:
                nxt = lines[i]
                if nxt.strip() == "" or re.match(r"^(\s*)-\s+", nxt) \
                        or re.match(r"^#{1,6}\s+", nxt.strip()) or is_table_line(nxt):
                    break
                buf.append(nxt.strip())
                i += 1
            blocks.append({"type": "li", "depth": depth, "text": " ".join(buf)})
            continue

        if stripped.startswith(">"):
            buf = [stripped.lstrip(">").strip()]
            i += 1
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append({"type": "p", "text": " ".join(buf)})
            continue

        # plain paragraph: consume until blank line / heading / list / table / hr
        buf = [stripped]
        i += 1
        while i < n:
            nxt = lines[i]
            if nxt.strip() == "" or re.match(r"^(\s*)-\s+", nxt) \
                    or re.match(r"^#{1,6}\s+", nxt.strip()) or is_table_line(nxt) \
                    or nxt.strip() == "---":
                break
            buf.append(nxt.strip())
            i += 1
        blocks.append({"type": "p", "text": " ".join(buf)})

    return blocks


def block_key(b):
    if b["type"] == "table":
        return "TABLE::" + b["text"]
    if b["type"] == "h":
        return f"H{b['level']}::{b['text']}"
    if b["type"] == "li":
        return f"LI{b['depth']}::{b['text']}"
    return "P::" + b["text"]


# ---------- inline formatting parser ----------

def parse_inline(text):
    """Returns list of (text, flags). '`code`', '**bold**', '_italic_'/
    '*italic*', and '~strike~'/'~~strike~~' spans all require a matching
    closing delimiter later in the text; an unpaired delimiter (a lone
    underscore in a word like "file_name", a stray asterisk or tilde) is
    left as a literal character instead of toggling formatting on for the
    rest of the block. Delimiters nest correctly (e.g. "**bold _and
    italic_**"), but content is not recursively re-parsed across different
    delimiter *spans* placed side by side -- only true nesting (one
    delimiter opened inside another, closed before the outer one) combines."""
    return _parse_inline(text, bold=False, italic=False, strike=False)


def _parse_inline(text, bold, italic, strike):
    runs = []
    buf = ""
    i = 0
    n = len(text)

    def flush():
        nonlocal buf
        if buf:
            runs.append((buf, {"bold": bold, "italic": italic, "strike": strike}))
            buf = ""

    while i < n:
        if text[i] == "`":
            end = text.find("`", i + 1)
            if end != -1:
                flush()
                runs.append((text[i + 1:end], {"bold": False, "italic": False, "strike": False}))
                i = end + 1
                continue
        if text[i:i + 2] == "**":
            end = text.find("**", i + 2)
            if end != -1:
                flush()
                runs.extend(_parse_inline(text[i + 2:end], True, italic, strike))
                i = end + 2
                continue
        if text[i] == "~":
            j = i
            while j < n and text[j] == "~":
                j += 1
            close_start = text.find("~", j)
            if close_start != -1:
                close_end = close_start
                while close_end < n and text[close_end] == "~":
                    close_end += 1
                flush()
                runs.extend(_parse_inline(text[j:close_start], bold, italic, True))
                i = close_end
                continue
        if text[i] in "_*":
            delim = text[i]
            end = text.find(delim, i + 1)
            if end != -1:
                flush()
                runs.extend(_parse_inline(text[i + 1:end], bold, True, strike))
                i = end + 1
                continue
        buf += text[i]
        i += 1
    flush()
    return runs


_WORD_RE = re.compile(r"\S+|\s+")


def inline_to_tokens(text):
    """Flatten parse_inline() output into (word_or_space, flags) tokens."""
    tokens = []
    for run_text, flags in parse_inline(text):
        for piece in _WORD_RE.findall(run_text):
            tokens.append((piece, flags))
    return tokens


# ---------- OOXML helpers ----------

class Redliner:
    def __init__(self, author, date):
        self.author = author
        self.date = date

    def next_id(self):
        return str(next(_rev_id))

    def _set_rpr(self, r_el, flags):
        if not any(flags.values()):
            return
        rpr = OxmlElement("w:rPr")
        if flags.get("bold"):
            rpr.append(OxmlElement("w:b"))
        if flags.get("italic"):
            rpr.append(OxmlElement("w:i"))
        if flags.get("strike"):
            rpr.append(OxmlElement("w:strike"))
        r_el.insert(0, rpr)

    def append_plain_run(self, paragraph, text, flags):
        if text == "":
            return
        run = paragraph.add_run(text)
        run.bold = True if flags.get("bold") else None
        run.italic = True if flags.get("italic") else None
        if flags.get("strike"):
            run.font.strike = True

    def append_tracked_run(self, paragraph, text, action, flags):
        if text == "":
            return
        wrap = OxmlElement("w:ins" if action == "ins" else "w:del")
        wrap.set(qn("w:id"), self.next_id())
        wrap.set(qn("w:author"), self.author)
        wrap.set(qn("w:date"), self.date)
        r = OxmlElement("w:r")
        self._set_rpr(r, flags)
        t = OxmlElement("w:delText" if action == "del" else "w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        wrap.append(r)
        paragraph._p.append(wrap)

    def render_tokens(self, paragraph, tokens, action=None):
        for word, flags in tokens:
            if action is None:
                self.append_plain_run(paragraph, word, flags)
            else:
                self.append_tracked_run(paragraph, word, action, flags)

    def render_word_diff(self, paragraph, old_text, new_text):
        old_tok = inline_to_tokens(old_text)
        new_tok = inline_to_tokens(new_text)
        sm = SequenceMatcher(None, [w for w, _ in old_tok], [w for w, _ in new_tok], autojunk=False)
        for tag, i1, i2, j1, j2 in sm.get_opcodes():
            if tag == "equal":
                self.render_tokens(paragraph, new_tok[j1:j2], None)
            elif tag == "delete":
                self.render_tokens(paragraph, old_tok[i1:i2], "del")
            elif tag == "insert":
                self.render_tokens(paragraph, new_tok[j1:j2], "ins")
            elif tag == "replace":
                self.render_tokens(paragraph, old_tok[i1:i2], "del")
                self.render_tokens(paragraph, new_tok[j1:j2], "ins")

    def mark_row_tracked(self, tc, action):
        tr = tc.getparent()
        trPr = tr.find(qn("w:trPr"))
        if trPr is None:
            trPr = OxmlElement("w:trPr")
            tr.insert(0, trPr)
        tag = "w:ins" if action == "ins" else "w:del"
        if trPr.find(qn(tag)) is None:
            mark = OxmlElement(tag)
            mark.set(qn("w:id"), self.next_id())
            mark.set(qn("w:author"), self.author)
            mark.set(qn("w:date"), self.date)
            trPr.append(mark)


STYLE_FOR_DEPTH = {0: "List Bullet", 1: "List Bullet 2", 2: "List Bullet 3"}


def new_paragraph_for_block(doc, block):
    if block["type"] == "h":
        level = min(block["level"], 9)
        return doc.add_paragraph(style=f"Heading {level}")
    if block["type"] == "li":
        style = STYLE_FOR_DEPTH.get(block["depth"], "List Bullet 3")
        return doc.add_paragraph(style=style)
    if block["type"] == "hr":
        p = doc.add_paragraph()
        p.alignment = 1
        return p
    return doc.add_paragraph()


def render_table(doc, rl, rows, inserted_row_idx=None, deleted_all=False):
    if not rows:
        return
    ncols = len(rows[0])
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    for r_i, row in enumerate(rows):
        row_cells = table.add_row().cells
        mark_whole_row_ins = (inserted_row_idx == "ALL") or (inserted_row_idx == r_i)
        if len(row) != ncols:
            raise ValueError(
                f"table row {r_i} has {len(row)} cell(s), expected {ncols} "
                f"(from the header row) -- fix the ragged markdown table "
                f"before regenerating: {row!r}")
        for c_i, cell_text in enumerate(row):
            cell = row_cells[c_i]
            p = cell.paragraphs[0]
            toks = inline_to_tokens(cell_text)
            if r_i == 0:
                toks = [(w, {**f, "bold": True}) for w, f in toks]
            if deleted_all:
                rl.render_tokens(p, toks, "del")
                rl.mark_row_tracked(cell._tc, "del")
            elif mark_whole_row_ins:
                rl.render_tokens(p, toks, "ins")
                rl.mark_row_tracked(cell._tc, "ins")
            else:
                rl.render_tokens(p, toks, None)


def render_table_row_diff(doc, rl, old_rows, new_rows):
    old_keys = ["|".join(r) for r in old_rows]
    new_keys = ["|".join(r) for r in new_rows]
    sm = SequenceMatcher(None, old_keys, new_keys, autojunk=False)
    old_ncols = len(old_rows[0]) if old_rows else 0
    new_ncols = len(new_rows[0]) if new_rows else 0
    # old and new table versions may legitimately differ in column count (a
    # column added/removed is a valid edit, not raggedness); render at the
    # wider of the two and pad the narrower side's rows.
    ncols = max(old_ncols, new_ncols)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    row_i = 0

    def add_row(cells_text, action, expected_ncols):
        nonlocal row_i
        if len(cells_text) != expected_ncols:
            raise ValueError(
                f"table row {row_i} has {len(cells_text)} cell(s), expected "
                f"{expected_ncols} (from its table's own header row) -- fix "
                f"the ragged markdown table before regenerating: {cells_text!r}")
        row_cells = table.add_row().cells
        padded = list(cells_text) + [""] * (ncols - len(cells_text))
        for c_i, cell_text in enumerate(padded):
            cell = row_cells[c_i]
            p = cell.paragraphs[0]
            toks = inline_to_tokens(cell_text)
            if row_i == 0:
                toks = [(w, {**f, "bold": True}) for w, f in toks]
            if action is None:
                rl.render_tokens(p, toks, None)
            else:
                rl.render_tokens(p, toks, action)
                rl.mark_row_tracked(cell._tc, action)
        row_i += 1

    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            for k in range(i1, i2):
                add_row(new_rows[j1 + (k - i1)], None, new_ncols)
        elif tag == "delete":
            for k in range(i1, i2):
                add_row(old_rows[k], "del", old_ncols)
        elif tag == "insert":
            for k in range(j1, j2):
                add_row(new_rows[k], "ins", new_ncols)
        elif tag == "replace":
            for k in range(i1, i2):
                add_row(old_rows[k], "del", old_ncols)
            for k in range(j1, j2):
                add_row(new_rows[k], "ins", new_ncols)


def render_block_plain(doc, rl, block):
    p = new_paragraph_for_block(doc, block)
    if block["type"] == "hr":
        p.add_run("\u2014" * 24)
        return
    if block["type"] == "table":
        render_table(doc, rl, block["rows"], inserted_row_idx=None)
        return
    rl.render_tokens(p, inline_to_tokens(block["text"]), None)


def render_block_tracked(doc, rl, block, action):
    p = new_paragraph_for_block(doc, block)
    if block["type"] == "hr":
        rl.append_tracked_run(p, "—" * 24, action,
                               {"bold": False, "italic": False, "strike": False})
        return
    if block["type"] == "table":
        render_table(doc, rl, block["rows"], inserted_row_idx="ALL" if action == "ins" else None,
                      deleted_all=(action == "del"))
        return
    rl.render_tokens(p, inline_to_tokens(block["text"]), action)


# ---------- top-level block diff / render ----------

def build(old_blocks, new_blocks, doc, rl):
    old_keys = [block_key(b) for b in old_blocks]
    new_keys = [block_key(b) for b in new_blocks]
    sm = SequenceMatcher(None, old_keys, new_keys, autojunk=False)

    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            for k in range(j1, j2):
                render_block_plain(doc, rl, new_blocks[k])
        elif tag == "delete":
            for k in range(i1, i2):
                render_block_tracked(doc, rl, old_blocks[k], "del")
        elif tag == "insert":
            for k in range(j1, j2):
                render_block_tracked(doc, rl, new_blocks[k], "ins")
        elif tag == "replace":
            old_sub = old_blocks[i1:i2]
            new_sub = new_blocks[j1:j2]
            both_one = len(old_sub) == 1 and len(new_sub) == 1
            if both_one and old_sub[0]["type"] == new_sub[0]["type"] == "table":
                render_table_row_diff(doc, rl, old_sub[0]["rows"], new_sub[0]["rows"])
            elif both_one and old_sub[0]["type"] == new_sub[0]["type"]:
                ob, nb = old_sub[0], new_sub[0]
                p = new_paragraph_for_block(doc, nb)
                rl.render_word_diff(p, ob["text"], nb["text"])
            else:
                for ob in old_sub:
                    render_block_tracked(doc, rl, ob, "del")
                for nb in new_sub:
                    render_block_tracked(doc, rl, nb, "ins")


def build_redline(old_text, new_text, author=DEFAULT_AUTHOR, date=None):
    """Returns a docx.Document with tracked changes from old_text to new_text."""
    if not author:
        raise ValueError(
            "author is required (pass author=... or set REDLINE_AUTHOR) -- "
            "this tool has no default attribution. A generic placeholder or "
            "someone else's org name silently committed to a document is "
            "worse than an explicit error.")
    date = date or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    rl = Redliner(author, date)
    doc = Document()
    build(parse_blocks(old_text), parse_blocks(new_text), doc, rl)
    return doc


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("old", help="path to the original Markdown file")
    ap.add_argument("new", help="path to the revised Markdown file")
    ap.add_argument("out", help="output .docx path")
    ap.add_argument("--author", default=DEFAULT_AUTHOR,
                     help="name attributed to the tracked changes. Required: pass "
                          "this or set REDLINE_AUTHOR. No default is provided on "
                          "purpose -- this is a shared public tool, not tied to "
                          "any one org or agent, so it will not guess whose name "
                          "belongs in your reviewers' document.")
    ap.add_argument("--date", default=None,
                     help="ISO-8601 UTC timestamp for the revisions (default: now)")
    args = ap.parse_args()
    if not args.author:
        ap.error("--author is required (or set the REDLINE_AUTHOR environment "
                  "variable) -- this tool has no default attribution.")

    old_text = open(args.old, encoding="utf-8").read()
    new_text = open(args.new, encoding="utf-8").read()
    doc = build_redline(old_text, new_text, args.author, args.date)
    doc.save(args.out)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
